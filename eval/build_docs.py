#!/usr/bin/env python3
"""由 eval/synth.js 產生的合成劇本，輸出 docx 與 PDF 測試檔（附標準答案）。

用法：
    node eval/synth.js eval/out/synth
    python3 eval/build_docs.py eval/out/synth eval/out/docs
需要：pip install python-docx reportlab；系統有 WenQuanYi Zen Hei 字型（或用 --font 指定）。

PDF 的頁眉、頁尾、頁碼是畫在頁面邊緣的（真實 PDF 的情況），不是內文行；
標準答案裡也把它們依閱讀順序列為 N 行，供評測時對齊（被前處理移除的 N 行不算錯）。
"""
import argparse
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import CIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

DEFAULT_FONT = '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'


def load(indir, vid):
    with open(os.path.join(indir, vid + '.gold.json'), encoding='utf-8') as f:
        return json.load(f)


def save_gold(outdir, vid, desc, roles, gold, partial=False):
    # partial：只用了原劇本的一部分（塞不進版面的行略過），端到端對照原劇本沒有意義，評測只看行級
    for i, g in enumerate(gold):
        g['n'] = i + 1
    out = {'id': vid, 'desc': desc, 'roles': roles, 'gold': gold}
    if partial:
        out['partial'] = True
    with open(os.path.join(outdir, vid + '.gold.json'), 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)


def make_docx(indir, outdir, src, vid, desc, centered_label=None):
    d = load(indir, src)
    doc = Document()
    for g in d['gold']:
        p = doc.add_paragraph()
        r = p.add_run(g['text'])
        if g['label'] == 'H':
            r.bold = True
        if centered_label and g['label'] == centered_label:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(os.path.join(outdir, vid + '.docx'))
    save_gold(outdir, vid, desc, d['roles'], d['gold'])


# 測試字型（文泉驛正黑）沒有的字，換成字型有的近似字；PDF 與標準答案同步替換，確保比對的是抽取能力而非缺字
SUBST = {'\U00028cd2': '屌', '\u22ef': '…', '\U000269f2': '脷', '\U00028d99': '隻'}


def subst(s):
    for a, b in SUBST.items():
        s = s.replace(a, b)
    return s


def make_pdf(indir, outdir, src, vid, desc, header_fn, footer_fn, per_page=36, font=DEFAULT_FONT):
    d = load(indir, src)
    pdfmetrics.registerFont(TTFont('Zen', font, subfontIndex=0))
    W, H = 900, 842
    c = canvas.Canvas(os.path.join(outdir, vid + '.pdf'), pagesize=(W, H))
    gold = []
    body = [dict(g, text=subst(g['text'])) for g in d['gold']]
    face = pdfmetrics.getFont('Zen').face
    missing = {ch for g in body for ch in g['text'] if ord(ch) not in face.charToGlyph}
    if missing:
        print('警告：字型缺字，請加進 SUBST：', ''.join(sorted(missing)))
    pages = [body[i:i + per_page] for i in range(0, len(body), per_page)]
    for pi, rows in enumerate(pages, 1):
        h = header_fn(pi)
        f = footer_fn(pi)
        if h:
            c.setFont('Zen', 9)
            c.drawString(60, H - 36, h)
            gold.append({'text': h, 'label': 'N', 'role': ''})
        c.setFont('Zen', 11)
        y = H - 80
        for g in rows:
            c.drawString(60, y, g['text'])
            gold.append({'text': g['text'], 'label': g['label'], 'role': g['role']})
            y -= 17
        if f:
            c.setFont('Zen', 9)
            c.drawString(W / 2 - 20, 30, f)
            gold.append({'text': f, 'label': 'N', 'role': ''})
        c.showPage()
    c.save()
    save_gold(outdir, vid, desc, d['roles'], gold)


def make_pdf_2col(indir, outdir, src, vid, desc, font=DEFAULT_FONT):
    """書籍式雙欄 PDF（得獎劇本集的版面）：每頁兩欄、左欄讀完才讀右欄；欄內「名字｜台詞」兩個儲存格、沒有冒號；
    兩欄的基線對齊（pdf.js 會把左右欄同一水平線的字併成一行）；頁尾奇數頁靠右、偶數頁靠左
    （左半邊的頁尾若不先分出去，會卡在左欄與右欄之間）。標準答案依閱讀順序：左欄、右欄、頁尾。
    塞不進欄寬的行略過（partial）；場次標題前空一行（欄頂／欄底會出現 2 倍行距）。"""
    d = load(indir, src)
    pdfmetrics.registerFont(TTFont('Zen', font, subfontIndex=0))
    W, H = 1000, 842
    FS, PITCH, TOP, ROWS = 10, 14, 776, 51
    X0 = (57, 540)                  # 兩欄名字儲存格的 x：相差 483，欄間空白 477–540
    CELL, MAXW = 60, 360            # 台詞儲存格在名字之後 60pt（容得下「偉／玲／朗」）、寬 360pt
    wid = lambda t, fs=FS: pdfmetrics.stringWidth(t, 'Zen', fs)
    body = [dict(g, text=subst(g['text'])) for g in d['gold']]
    face = pdfmetrics.getFont('Zen').face
    missing = {ch for g in body for ch in g['text'] if ord(ch) not in face.charToGlyph}
    if missing:
        print('警告：字型缺字，請加進 SUBST：', ''.join(sorted(missing)))
    # 排成一欄一欄的列：(kind, name, text, goldrow)；kind: S 名字＋台詞、C 台詞續行、D 指示、H 標題、_ 空行
    seq, skipped = [], 0
    for g in body:
        t = g['text']
        if g['label'] == 'S' and '：' in t:
            name, rest = t.split('：', 1)
            if wid(rest) > MAXW:
                skipped += 1
                continue
            seq.append(('S', name, rest, g))
        elif g['label'] == 'C':
            if wid(t) > MAXW:
                skipped += 1
                continue
            seq.append(('C', '', t, g))
        elif g['label'] in ('D', 'H'):
            if wid(t, 12 if g['label'] == 'H' else FS) > CELL + MAXW:
                skipped += 1
                continue
            if g['label'] == 'H':
                seq.append(('_', '', '', None))
            seq.append((g['label'], '', t, g))
    cols, cur, used = [], [], 0
    for item in seq:
        need = 1
        if item[0] == '_' and not cur:
            continue                                      # 欄頂不留空行
        if used + need > ROWS:
            cols.append(cur); cur, used = [], 0
            if item[0] == '_':
                continue
        cur.append(item); used += need
    if cur:
        cols.append(cur)
    c = canvas.Canvas(os.path.join(outdir, vid + '.pdf'), pagesize=(W, H))
    gold = []
    for pi in range(0, len(cols), 2):
        pno = pi // 2 + 1
        for side, col in enumerate(cols[pi:pi + 2]):
            x0 = X0[side]
            for row, (kind, name, text, g) in enumerate(col):
                y = TOP - row * PITCH
                if kind == '_':
                    continue
                if kind == 'S':
                    c.setFont('Zen', FS); c.drawString(x0, y, name); c.drawString(x0 + CELL, y, text)
                elif kind == 'C':
                    c.setFont('Zen', FS); c.drawString(x0 + CELL, y, text)
                elif kind == 'H':
                    c.setFont('Zen', 12); c.drawString(x0, y, text)
                else:
                    c.setFont('Zen', FS); c.drawString(x0, y, text)
                gold.append({'text': g['text'], 'label': g['label'], 'role': g['role']})
        c.setFont('Zen', 8)
        if pno % 2:
            c.drawString(W - 200, 23, '合成劇本選集'); c.drawString(W - 48, 24, str(pno))
            foot = '合成劇本選集 %d' % pno
        else:
            c.drawString(35, 24, str(pno)); c.drawString(65, 24, '離地，到着 劇本集')
            foot = '%d 離地，到着 劇本集' % pno
        gold.append({'text': foot, 'label': 'N', 'role': ''})
        c.showPage()
    c.save()
    save_gold(outdir, vid, desc, d['roles'], gold, partial=True)
    print('%s：%d 頁、%d 欄；略過塞不進欄寬的 %d 行' % (vid, (len(cols) + 1) // 2, len(cols), skipped))


# ---- 不內嵌字型、靠預先定義 CMap（Adobe-CNS1／ETen-B5）的繁中 PDF：橫排與直排 ----
# 真實的繁中劇本 PDF 常是這種；pdf.js 沒帶 cMapUrl 時會把整段中文丟掉（抽出 0 行）。
# 內容是自寫的通用對話（只用 Big5 能編碼的字）。
_CN = '一二三四五六七八九'


def _cns1_lines():
    out, n = [], 0
    for rnd in range(3):                                  # 場次與結尾指示都不重複，避免被當成頁眉頁尾
        for title, body in [
            ('客廳', [('D', '', '（兄趴在地上，妹坐在沙發上看電視。）'), ('S', '兄', '兄：好好喔！'), ('S', '妹', '妹：但其實我不太想吃。'),
                     ('S', '兄', '兄：為什麼？你很討厭耶。'), ('S', '妹', '妹：你又不是不知道我已經吃了很多天，'), ('C', '', '每天都是一樣的東西。'),
                     ('D', '', '（兄起身走到門邊。）'), ('S', '兄', '兄：那我出去走走。'), ('S', '妹', '妹：路上小心。')]),
            ('廚房', [('D', '', '（燈亮，妹在整理桌面。）'), ('S', '兄', '兄：我回來了。'), ('S', '妹', '妹：外面怎麼樣？'),
                     ('S', '兄', '兄：很冷，而且下著雨，我淋得全身都濕了，'), ('C', '', '還差點被車撞到。'), ('S', '妹', '妹：下次記得帶傘。'),
                     ('D', '', '（兩人沉默。）'), ('S', '兄', '兄：好。'), ('S', '妹', '妹：晚安。')]),
            ('房間', [('D', '', '（黑暗中傳來腳步聲。）'), ('S', '妹', '妹：你還沒睡嗎？'), ('S', '兄', '兄：睡不著。'),
                     ('S', '妹', '妹：那我們聊聊天吧。'), ('S', '兄', '兄：好啊。')]),
        ]:
            n += 1
            out.append(('H', '', '第%s場　%s' % (_CN[n - 1], title)))
            out.extend(body)
            out.append(('D', '', '（第%s段結束。）' % _CN[n - 1]))
    return out


CNS1_LINES = _cns1_lines()


def make_cns1_pdf(outdir, vertical, per_page=24):
    if any(True for _, _, t in CNS1_LINES if t.encode('big5') is None):
        raise SystemExit('big5')
    face = 'MSung-Light-ETenms-B5-' + ('V' if vertical else 'H')
    pdfmetrics.registerFont(CIDFont('MSung-Light', 'ETenms-B5-' + ('V' if vertical else 'H')))
    enc = lambda t: t.encode('big5').decode('latin-1')
    vid = 'pdf-cns1-' + ('v' if vertical else 'h')
    c = canvas.Canvas(os.path.join(outdir, vid + '.pdf'), pagesize=(612, 792))
    gold = []
    pages = [CNS1_LINES[i:i + per_page] for i in range(0, len(CNS1_LINES), per_page)]
    for pi, rows in enumerate(pages, 1):
        if vertical:
            c.setFont(face, 14)
            x = 560
            for lab, role, t in rows:                      # 每個欄位一行，欄由右到左
                c.drawString(x, 720, enc(t)); x -= 22
                gold.append({'text': t, 'label': lab, 'role': role})
            c.setFont(face, 9); c.drawString(40, 60, enc('第 %d 頁' % pi))   # 頁尾頁碼（橫排）
            gold.append({'text': '第 %d 頁' % pi, 'label': 'N', 'role': ''})
        else:
            c.setFont(face, 9); c.drawString(60, 760, enc('通用示範劇本')); gold.append({'text': '通用示範劇本', 'label': 'N', 'role': ''})
            c.setFont(face, 12); y = 720
            for lab, role, t in rows:
                c.drawString(60, y, enc(t)); y -= 22
                gold.append({'text': t, 'label': lab, 'role': role})
            c.setFont(face, 9); c.drawString(290, 40, enc('第 %d 頁' % pi)); gold.append({'text': '第 %d 頁' % pi, 'label': 'N', 'role': ''})
        c.showPage()
    c.save()
    save_gold(outdir, vid, 'PDF：不內嵌字型的繁中（ETen-B5）' + ('直排，欄由右到左' if vertical else '橫排'),
              [{'id': '兄', 'name': '兄'}, {'id': '妹', 'name': '妹'}], gold)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('indir', nargs='?', default='eval/out/synth')
    ap.add_argument('outdir', nargs='?', default='eval/out/docs')
    ap.add_argument('--font', default=DEFAULT_FONT)
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)

    make_docx(a.indir, a.outdir, 'colon-fw', 'docx-colon', 'docx：角色簡稱＋全形冒號')
    make_docx(a.indir, a.outdir, 'centered', 'docx-centered', 'docx：角色名置中獨立成行', centered_label='N')
    make_pdf(a.indir, a.outdir, 'wrapped', 'pdf-wrapped', 'PDF：折行台詞，每頁頁眉＋頁尾「- n -」',
             lambda p: '《離地，到着》劇本草稿 v3', lambda p: '- %d -' % p, font=a.font)
    make_pdf(a.indir, a.outdir, 'centered-wrapped', 'pdf-centered', 'PDF：置中角色名，奇偶頁不同頁眉，頁尾「第 n 頁」',
             lambda p: '《離地，到着》· 單數頁' if p % 2 else '劇本草稿 v3 · 雙數頁', lambda p: '第 %d 頁' % p, font=a.font)
    make_pdf(a.indir, a.outdir, 'noisy', 'pdf-messy', 'PDF：封面、目錄、內文頁碼雜訊再加頁眉頁尾',
             lambda p: '《離地，到着》機密　內部排練用', lambda p: 'Page %d of 99' % p, per_page=40, font=a.font)
    make_pdf_2col(a.indir, a.outdir, 'wrapped', 'pdf-2col', 'PDF：書籍式雙欄（名字｜台詞儲存格、沒有冒號），奇數頁頁尾靠右、偶數頁頁尾靠左', font=a.font)
    make_cns1_pdf(a.outdir, False)
    make_cns1_pdf(a.outdir, True)
    print('已輸出到', a.outdir, sorted(os.listdir(a.outdir)))


if __name__ == '__main__':
    main()
